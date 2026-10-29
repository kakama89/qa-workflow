#!/usr/bin/env python3
"""Write approved test cases into a copy of a test-case XLSX template.

The template layout (Test Case sheet, header row, columns, ID convention) is
detected the same way as validate_workbook.py, so team-specific templates
work without code changes. The template file itself is never modified.

Usage:
  python export_workbook.py --template TEMPLATE.xlsx --cases cases.json --output OUT.xlsx
  python export_workbook.py --profile NAME --cases cases.json --output OUT.xlsx

--profile loads a tester profile saved with qa_profiles.py (template, layout,
metadata); options given explicitly override the profile.

cases.json (UTF-8):
  {
    "metadata": {"Project Name:": "...", "Cover!E16": "1.0"},
    "cases": [
      {"section": "Login"},
      {"id": "", "title": "...", "precondition": "...", "steps": "1. ...\\n2. ...",
       "data": "...", "expected": "...", "priority": "High",
       "columns": {"Any other header text": "value"}}
    ]
  }
Keys are column roles (see validate_workbook.py --help). Leave "id" empty to
use the template's ID formula. "metadata" fills single fields by label text
(the value goes into the first cell right of the label) or by "Sheet!A1".
Run "validate_workbook.py validate" afterwards.

Exit codes: 0 = written, 1 = refused (data problems), 2 = usage or file error.
"""
from __future__ import annotations

import argparse
import json
import re
import sys
from copy import copy
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))
from validate_workbook import (  # noqa: E402
    HEADER_ALIASES, NOT_EXECUTED, RANGE_RE, UsageError, compress_rows, formula_text,
    is_blank, locate, norm, open_workbook, secret_kinds,
)
from openpyxl.formula.translate import Translator  # noqa: E402
from openpyxl.styles import Font  # noqa: E402
from openpyxl.utils import column_index_from_string, get_column_letter  # noqa: E402
from openpyxl.utils.cell import range_boundaries  # noqa: E402


def load_cases(path):
    try:
        data = json.loads(Path(path).read_text(encoding="utf-8-sig"))
    except FileNotFoundError:
        raise UsageError(f"Cases file not found: {path}")
    except json.JSONDecodeError as exc:
        raise UsageError(f"Cases file is not valid JSON: {exc}")
    cases = data.get("cases") if isinstance(data, dict) else data
    if not isinstance(cases, list) or not cases:
        raise UsageError('Cases file must contain a non-empty "cases" list.')
    metadata = data.get("metadata") if isinstance(data, dict) else None
    if metadata is not None and not isinstance(metadata, dict):
        raise UsageError('"metadata" must be an object of {"label or Sheet!A1": value}.')
    return cases, metadata or {}


def label_key(text) -> str:
    return norm(text).rstrip(":").strip()


def fill_metadata(wb, metadata, notes):
    """Write metadata values next to their label cells (or into Sheet!A1 cells)."""
    for key, value in metadata.items():
        m = re.fullmatch(r"(?:'((?:[^']|'')+)'|([^!]+))!\$?([A-Za-z]{1,3})\$?(\d+)", key.strip())
        if m:
            name = (m.group(1) or m.group(2)).replace("''", "'")
            if name not in wb.sheetnames:
                notes.append(f"WARNING metadata {key!r}: no sheet {name!r}; skipped")
                continue
            wb[name][f"{m.group(3).upper()}{m.group(4)}"].value = value
            notes.append(f"metadata {key} filled")
            continue
        hits = [(ws, cell) for ws in wb.worksheets for row in ws.iter_rows() for cell in row
                if isinstance(cell.value, str) and label_key(cell.value) == label_key(key)]
        if not hits:
            notes.append(f"WARNING metadata label {key!r} not found; skipped")
            continue
        ws, label = hits[0]
        col = label.column + 1
        for rng in ws.merged_cells.ranges:
            if label.coordinate in rng:
                col = rng.max_col + 1
        target = ws.cell(label.row, col)
        for rng in ws.merged_cells.ranges:
            if target.coordinate in rng:
                target = ws.cell(rng.min_row, rng.min_col)
        target.value = value
        extra = f" (label also appears {len(hits) - 1} more time(s); use Sheet!A1 to pick another)" if len(hits) > 1 else ""
        notes.append(f"metadata {key!r} -> {ws.title}!{target.coordinate}{extra}")


def check_cases(cases, cols, header_cols, stage):
    """Return a list of problems; nothing is written when any exist."""
    problems = []
    seen_ids = {}
    for i, case in enumerate(cases, start=1):
        if not isinstance(case, dict):
            problems.append(f"case #{i}: must be an object")
            continue
        if "section" in case:
            continue
        if not is_blank(case.get("id")):
            key = norm(case["id"])
            if key in seen_ids:
                problems.append(f"case #{i}: duplicate id {case['id']!r} (also case #{seen_ids[key]})")
            seen_ids.setdefault(key, i)
        for key in case:
            if key == "columns":
                for header in case["columns"]:
                    if norm(header) not in header_cols:
                        problems.append(f"case #{i}: template has no column {header!r}")
            elif key not in HEADER_ALIASES:
                problems.append(f"case #{i}: unknown key {key!r}")
            elif key not in cols and not is_blank(case[key]):
                problems.append(f"case #{i}: template has no {key!r} column (map it with --col {key}=<column>)")
        if stage == "design" and norm(case.get("result")) not in NOT_EXECUTED:
            problems.append(f"case #{i}: result {case.get('result')!r} is an execution result; "
                            "design exports must leave it blank or Untested")
        values = [case.get(k) for k in HEADER_ALIASES] + list((case.get("columns") or {}).values())
        for v in values:
            if isinstance(v, str) and secret_kinds(v):
                problems.append(f"case #{i}: looks like it contains a credential; use a placeholder such as <TOKEN>")
                break
    return problems


def find_id_formula(ws, id_col, start):
    for r in range(start, ws.max_row + 1):
        f = formula_text(ws.cell(r, id_col).value)
        if f:
            return r, f
    return None


def clear_rows(ws, first, last, max_col, keep_col=None):
    for rng in [r for r in ws.merged_cells.ranges if r.max_row >= first and r.min_row <= last]:
        ws.unmerge_cells(str(rng))
    for r in range(first, last + 1):
        for c in range(1, max_col + 1):
            if c != keep_col:
                ws.cell(r, c).value = None


def extend_ranges(wb, layout, last_row, notes):
    """Stretch dropdowns, summary formulas and filters that stop above last_row."""
    ws = layout["sheet"].ws
    cols = set(layout["cols"].values())
    head = layout["header_row"]
    first = layout["data_start"]

    for dv in ws.data_validations.dataValidation:
        new_parts, changed = [], False
        for rng in str(dv.sqref).split():
            c1, r1, c2, r2 = range_boundaries(rng)
            if r2 is not None and r2 < last_row and r2 > head and any((c1 or 1) <= c <= (c2 or 16384) for c in cols):
                rng = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{last_row}"
                changed = True
            new_parts.append(rng)
        if changed:
            dv.sqref = type(dv.sqref)(" ".join(new_parts))
            notes.append(f"dropdown extended to {' '.join(new_parts)}")

    def stretch(match, home):
        sheet_name = (match.group(1) or "").replace("''", "'") or match.group(2) or home
        c1, r1 = column_index_from_string(match.group(3)), int(match.group(4))
        c2, r2 = column_index_from_string(match.group(5)), int(match.group(6))
        text = match.group(0)
        if sheet_name == ws.title and c1 == c2 and c1 in cols and head < r1 <= first + 1 and r2 < last_row:
            return text[: text.rfind(match.group(6))] + str(last_row)
        return text

    for sheet in wb.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                f = formula_text(cell.value)
                if not f or not isinstance(cell.value, str) or (sheet is ws and cell.row >= first):
                    continue
                new = RANGE_RE.sub(lambda m: stretch(m, sheet.title), f)
                if new != f:
                    cell.value = new
                    notes.append(f"formula {sheet.title}!{cell.coordinate}: {f} -> {new}")

    if ws.auto_filter.ref:
        c1, r1, c2, r2 = range_boundaries(ws.auto_filter.ref)
        if r2 is not None and r2 < last_row:
            ws.auto_filter.ref = f"{get_column_letter(c1)}{r1}:{get_column_letter(c2)}{last_row}"
            notes.append(f"filter extended to {ws.auto_filter.ref}")

    for cf in ws.conditional_formatting:
        for rng in str(cf.sqref).split():
            _, _, _, r2 = range_boundaries(rng)
            if r2 is not None and head < r2 < last_row:
                notes.append(f"WARNING conditional formatting {rng} stops before row {last_row}; extend it manually")


def run(args) -> int:
    template, output = Path(args.template), Path(args.output)
    if output.exists():
        if output.resolve() == template.resolve():
            raise UsageError("Output must be a new file, not the template.")
        if not args.overwrite_output:
            raise UsageError(f"Output {output} already exists; choose a new name or pass --overwrite-output.")
    if template.suffix.lower() == ".xlsm" and output.suffix.lower() != ".xlsm":
        raise UsageError("Macro-enabled template: the output must also use .xlsm.")
    cases, metadata = load_cases(args.cases)
    for item in args.meta or []:
        if "=" not in item:
            raise UsageError(f"--meta expects LABEL=VALUE, got {item!r}")
        key, value = item.split("=", 1)
        metadata.setdefault(key, value)  # cases.json wins over profile/CLI defaults
    wb = open_workbook(template, data_only=False)
    if template.suffix.lower() == ".xlsm":
        from openpyxl import load_workbook
        wb = load_workbook(template, keep_vba=True)
    layout = locate(wb, None, args)
    if layout is None:
        raise UsageError("Could not find the Test Case header in the template; use --sheet/--header-row/--col.")
    ws, cols = layout["sheet"].ws, layout["cols"]
    header_cols = {norm(ws.cell(layout["header_row"], c).value): c for c in range(1, ws.max_column + 1)
                   if not is_blank(ws.cell(layout["header_row"], c).value)}
    table_width = max(header_cols.values())

    problems = check_cases(cases, cols, header_cols, args.stage)
    if problems:
        print("Refused; nothing was written:")
        for p in problems:
            print(f"- {p}")
        return 1

    id_col = cols.get("id")
    id_formula = find_id_formula(ws, id_col, layout["data_start"]) if id_col else None
    existing = layout["tc_rows"] + layout["section_rows"] + layout["id_only_rows"]
    if existing and args.existing == "error":
        print(f"Refused: the template already has data in rows {compress_rows(existing)}. "
              "Pass --existing clear (replace sample rows) or --existing append.")
        return 1
    start = layout["data_start"]
    if existing and args.existing == "clear":
        clear_rows(ws, start, max(existing), table_width, keep_col=id_col if id_formula else None)
    elif existing:
        start = max(existing) + 1

    style_row = layout["tc_rows"][0] if layout["tc_rows"] else start
    styles = {c: copy(ws.cell(style_row, c)._style) for c in range(1, table_width + 1)}
    title_col = cols.get("title") or min(c for role, c in cols.items() if role != "id")

    r = start
    for case in cases:
        for c in range(1, table_width + 1):
            if not ws.cell(r, c).has_style:
                ws.cell(r, c)._style = copy(styles[c])
        if "section" in case:
            cell = ws.cell(r, title_col, case["section"])
            cell.font = Font(bold=True)
        else:
            for role, c in cols.items():
                if role != "id" and not is_blank(case.get(role)):
                    ws.cell(r, c, case[role])
            for header, value in (case.get("columns") or {}).items():
                ws.cell(r, header_cols[norm(header)], value)
            if id_col and not is_blank(case.get("id")):
                ws.cell(r, id_col, case["id"])
            elif id_col and not id_formula:
                print(f"Refused: case at row {r} has no id and the template has no ID formula to generate one.")
                return 1
        if id_formula and is_blank(case.get("id")):
            origin = f"{get_column_letter(id_col)}{id_formula[0]}"
            ws.cell(r, id_col).value = Translator(id_formula[1], origin=origin).translate_formula(
                f"{get_column_letter(id_col)}{r}")
        r += 1
    last_row = r - 1

    notes = []
    extend_ranges(wb, layout, last_row, notes)
    fill_metadata(wb, metadata, notes)
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    written = sum(1 for c in cases if "section" not in c)
    print(f"Wrote {written} test case(s) and {len(cases) - written} section row(s) to {output} "
          f"(sheet {ws.title}, rows {start}-{last_row}).")
    for n in notes:
        print(f"- {n}")
    print("Next: run validate_workbook.py validate on the output with the same --profile or --template.")
    return 0


def main(argv=None) -> int:
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8", errors="replace")
        except (AttributeError, ValueError):
            pass
    p = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    p.add_argument("--template", help="template .xlsx/.xlsm, never modified (default: the profile's template)")
    p.add_argument("--profile", metavar="NAME", help="tester profile from qa_profiles.py ('default' = default)")
    p.add_argument("--home", help="profile folder (see qa_profiles.py)")
    p.add_argument("--project-dir", help="project root for project-scope profiles")
    p.add_argument("--meta", action="append", metavar="LABEL=VALUE",
                   help='fill a field by label text ("Originator:=ngannt23") or cell ("Cover!E13=...")')
    p.add_argument("--cases", required=True, help="UTF-8 JSON file with approved test cases")
    p.add_argument("--output", required=True, help="new output file")
    p.add_argument("--overwrite-output", action="store_true", help="replace an existing output file")
    p.add_argument("--existing", choices=("error", "clear", "append"), default="error",
                   help="what to do with rows already in the template (default: refuse)")
    p.add_argument("--stage", choices=("design", "execution"), default="design")
    p.add_argument("--sheet")
    p.add_argument("--header-row", type=int)
    p.add_argument("--data-start-row", type=int)
    p.add_argument("--col", action="append", metavar="ROLE=COLUMN")
    args = p.parse_args(argv)
    try:
        if args.profile:
            from qa_profiles import apply_profile, load_profile

            profile = load_profile(args.profile, args.home, args.project_dir)
            apply_profile(args, profile)
            args.template = args.template or profile["template_path"]
        if not args.template:
            raise UsageError("Pass --template FILE or --profile NAME.")
        return run(args)
    except UsageError as exc:
        print(f"ERROR: {exc}", file=sys.stderr)
        return 2


if __name__ == "__main__":
    raise SystemExit(main())
